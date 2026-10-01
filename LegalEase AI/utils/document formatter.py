from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from utils.text_utils import (
    sanitize_text,
    split_terms,
)


def is_heading(line: str) -> bool:

    stripped = line.strip()

    if not stripped:
        return False

    if (
        len(stripped) <= 80
        and stripped.upper() == stripped
        and any(char.isalpha() for char in stripped)
    ):
        return True

    return stripped.startswith(
        (
            "1.",
            "2.",
            "3.",
            "4.",
            "5.",
            "6.",
            "7.",
            "8.",
            "9.",
        )
    )


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_bytes: bytes | None = None,
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # Logo

    if logo_bytes:

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                BytesIO(logo_bytes),
                width=Inches(1.3)
            )

        except Exception:
            pass

    # Title

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        doc_type.upper()
    )

    title_run.bold = True

    title_run.font.name = (
        "Times New Roman"
    )

    title_run.font.size = Pt(16)

    # Body

    cleaned_text = sanitize_text(text)

    for raw_line in cleaned_text.split("\n"):

        line = raw_line.strip()

        if not line:

            document.add_paragraph()

            continue

        if is_heading(line):

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_before = Pt(8)

            paragraph.paragraph_format.space_after = Pt(4)

            run = paragraph.add_run(line)

            run.bold = True

            run.font.name = "Times New Roman"

            run.font.size = Pt(12)

        else:

            paragraph = document.add_paragraph(line)

            paragraph.paragraph_format.space_after = Pt(5)

            paragraph.paragraph_format.line_spacing = 1.08

    # Terms table

    parsed_terms = split_terms(terms)

    if parsed_terms:

        document.add_paragraph()

        heading = document.add_paragraph()

        heading_run = heading.add_run(
            "Key Terms Provided by User"
        )

        heading_run.bold = True

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"

        for index, term in enumerate(
            parsed_terms,
            start=1
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)

            cells[1].text = term

    # Footer

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-generated draft. "
        "Review for the applicable jurisdiction before use."
    )

    footer_run.font.size = Pt(8)

    output = BytesIO()

    document.save(output)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        doc_type: str,
        logo_bytes: bytes | None = None,
    ):

        super().__init__()

        self.doc_type = doc_type

        self.logo_bytes = logo_bytes

        self.set_margins(
            18,
            20,
            18
        )

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        if self.logo_bytes:

            try:

                self.image(
                    self.logo_bytes,
                    x=90,
                    y=8,
                    w=30,
                )

                self.ln(18)

            except Exception:
                pass

        self.set_font(
            "Helvetica",
            "B",
            14
        )

        self.cell(
            0,
            8,
            self.doc_type.upper(),
            align="C",
        )

        self.ln(10)

    def footer(self):

        self.set_y(-12)

        self.set_font(
            "Helvetica",
            "",
            7
        )

        self.cell(
            0,
            5,
            "LegalEase - AI-generated draft - Review before use",
            align="C",
        )


def format_pdf(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_bytes: bytes | None = None,
) -> bytes:

    pdf = LegalEasePDF(
        doc_type=doc_type,
        logo_bytes=logo_bytes,
    )

    pdf.add_page()

    cleaned_text = sanitize_text(text)

    for raw_line in cleaned_text.split("\n"):

        line = raw_line.strip()

        if not line:

            pdf.ln(3)

            continue

        if is_heading(line):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                6,
                line
            )

            pdf.ln(1)

        else:

            pdf.set_font(
                "Helvetica",
                "",
                10
            )

            pdf.multi_cell(
                0,
                5.5,
                line
            )

            pdf.ln(1)

    parsed_terms = split_terms(terms)

    if parsed_terms:

        pdf.set_font(
            "Helvetica",
            "B",
            11
        )

        pdf.cell(
            0,
            7,
            "Key Terms Provided by User"
        )

        pdf.ln(7)

        for index, term in enumerate(
            parsed_terms,
            start=1
        ):

            pdf.set_font(
                "Helvetica",
                "",
                9.5
            )

            pdf.multi_cell(
                0,
                5,
                f"{index}. {term}"
            )

            pdf.ln(1)

    result = pdf.output()

    return bytes(result)